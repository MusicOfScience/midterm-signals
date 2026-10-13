#!/usr/bin/env python3
"""Certified-results entrypoint for the 2022 historical admission.

Predictive inputs and scoreability come from ``admit_historical_2022_strict``.
This final layer replaces the brittle Harvard Dataverse results download with
FEC's own Federal Elections 2022 workbook. The workbook is fetched only after
all predictive inputs and admission decisions have been constructed by the
underlying pipeline.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from collections import defaultdict

from openpyxl import load_workbook

import admit_historical_2022_strict as strict

impl = strict.impl
_STRICT_FETCH = impl.fetch
_RESULTS_REQUEST_LABEL = "MIT Election Data and Science Lab House 1976-2022"
FEC_RESULTS_URL = "https://www.fec.gov/documents/5676/federalelections2022.xlsx"
FEC_RESULTS_SHA256 = "cdb258ea23803e50d752bcea19faa39b3c201eb960094bcb556c3a38c2350897"
FEC_HOUSE_SHEET = "8. US House Results by State"


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _number(value) -> float:
    if value in (None, ""):
        return 0.0
    try:
        return float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return 0.0


def _party_family(value) -> str:
    p = _text(value).upper()
    if p in {"D", "DFL", "D/R"} or p.startswith(("D/", "D*", "DEM")):
        return "DEM"
    if p == "R" or p.startswith(("R/", "R*", "REP")):
        return "REP"
    return "OTHER"


def _district_number(value: str) -> str:
    raw = _text(value).upper()
    if "UNEXPIRED" in raw:
        return ""
    match = re.search(r"\d+", raw)
    if match:
        return str(int(match.group(0)))
    if raw in {"AL", "AT LARGE", "AT-LARGE", "AT LARGE DISTRICT"}:
        return "1"
    return ""


def _resolved_party(parties: set[str], *, state: str, district: str, candidate: str) -> str:
    """Resolve a candidate's certified fusion lines to the D/R control target."""
    major = parties & {"DEM", "REP"}
    if len(major) > 1:
        raise RuntimeError(
            f"FEC candidate appears on both D and R lines for {state}-{district}: "
            f"{candidate!r} parties={sorted(parties)}"
        )
    if major:
        return next(iter(major))
    return "OTHER"


def fec_house_winners_csv(workbook: bytes) -> bytes:
    """Convert the certified FEC House sheet to the minimal result schema.

    FEC can list one candidate on several ballot/fusion-party lines (for
    example, a Democratic winner can also appear on a Working Families line),
    and repeats the winner marker on those lines. We therefore collapse rows to
    a candidate within each district before selecting the winner. This prevents
    ballot lines from being misread as distinct winning candidates.

    Winner markers are preferred after candidate-level collapse. If a district
    has no explicit marker, certified runoff votes are used when present,
    otherwise certified general-election votes. Model Lab needs party control,
    not the candidate vote margin, so only one final winner row is emitted.
    """
    digest = hashlib.sha256(workbook).hexdigest()
    if digest != FEC_RESULTS_SHA256:
        raise RuntimeError(
            f"FEC Federal Elections 2022 workbook hash changed: {digest}; expected {FEC_RESULTS_SHA256}"
        )

    wb = load_workbook(io.BytesIO(workbook), read_only=True, data_only=True)
    if FEC_HOUSE_SHEET not in wb.sheetnames:
        raise RuntimeError(f"FEC workbook missing sheet {FEC_HOUSE_SHEET!r}: {wb.sheetnames}")
    ws = wb[FEC_HOUSE_SHEET]
    iterator = ws.iter_rows(values_only=True)
    try:
        raw_header = next(iterator)
    except StopIteration as exc:
        raise RuntimeError("FEC House result sheet is empty") from exc
    header = [_text(v) for v in raw_header]
    index = {name: i for i, name in enumerate(header) if name}
    required = {
        "STATE ABBREVIATION", "CANDIDATE NAME", "PARTY",
        "GENERAL VOTES", "GE WINNER INDICATOR",
    }
    missing = sorted(required - set(index))
    district_col = "D" if "D" in index else "DISTRICT" if "DISTRICT" in index else ""
    runoff_cols = [name for name in index if name.startswith("GE RUNOFF ELECTION VOTES")]
    if missing or not district_col or not runoff_cols:
        raise RuntimeError(
            f"unexpected FEC 2022 House schema; missing={missing}, district_col={district_col!r}, "
            f"runoff_cols={runoff_cols}, columns={header}"
        )
    runoff_col = runoff_cols[0]

    # District -> normalized candidate -> candidate-level certified record.
    groups: dict[tuple[str, str], dict[str, dict]] = defaultdict(dict)
    for values in iterator:
        def get(name):
            pos = index[name]
            return values[pos] if pos < len(values) else None

        state = _text(get("STATE ABBREVIATION")).upper()
        if state not in impl.FIPS:
            continue
        district = _district_number(get(district_col))
        if not district:
            # The six at-large states can have a blank district cell in FEC
            # layouts; they are unambiguously district 1 in the 118th map.
            if state in strict.admission.AT_LARGE_2022:
                district = "1"
            else:
                continue
        candidate = _text(get("CANDIDATE NAME"))
        raw_party = _text(get("PARTY"))
        if not candidate or not raw_party:
            continue
        key = impl.norm_name(candidate)
        if not key:
            continue
        record = groups[(state, district)].setdefault(key, {
            "candidate": candidate,
            "parties": set(),
            "general": 0.0,
            "runoff": 0.0,
            "winner": False,
        })
        record["parties"].add(_party_family(raw_party))
        record["general"] += _number(get("GENERAL VOTES"))
        record["runoff"] += _number(get(runoff_col))
        record["winner"] = record["winner"] or _text(get("GE WINNER INDICATOR")).upper() == "W"

    if len(groups) != 435:
        raise RuntimeError(f"FEC certified House sheet classified {len(groups)} districts, expected 435")

    output = io.StringIO()
    writer = csv.DictWriter(
        output,
        fieldnames=["year", "stage", "state_po", "district", "candidate", "party", "candidatevotes"],
        lineterminator="\n",
    )
    writer.writeheader()
    nonmajor = []
    for (state, district), candidates in sorted(groups.items(), key=lambda item: (item[0][0], int(item[0][1]))):
        rows = list(candidates.values())
        marked = [row for row in rows if row["winner"]]
        if len(marked) == 1:
            winner = marked[0]
        elif len(marked) > 1:
            raise RuntimeError(
                f"FEC marks multiple distinct general-election winners for {state}-{district}: "
                f"{[row['candidate'] for row in marked]}"
            )
        else:
            use_runoff = any(row["runoff"] > 0 for row in rows)
            vote_key = "runoff" if use_runoff else "general"
            winner = max(rows, key=lambda row: row[vote_key])

        party = _resolved_party(
            winner["parties"], state=state, district=district, candidate=winner["candidate"]
        )
        if party not in {"DEM", "REP"}:
            nonmajor.append(
                f"{state}-{district}:{winner['candidate']}:{sorted(winner['parties'])}"
            )
        writer.writerow({
            "year": 2022,
            "stage": "GEN",
            "state_po": state,
            "district": district,
            "candidate": winner["candidate"],
            "party": party,
            "candidatevotes": 1,
        })
    if nonmajor:
        raise RuntimeError(f"FEC House winners outside D/R target: {nonmajor}")
    return output.getvalue().encode("utf-8")


def certified_fetch(label: str, url: str, *args, **kwargs) -> bytes:
    if label != _RESULTS_REQUEST_LABEL:
        return _STRICT_FETCH(label, url, *args, **kwargs)
    workbook = _STRICT_FETCH(
        "FEC Federal Elections 2022 certified House results workbook",
        FEC_RESULTS_URL,
        timeout=600,
    )
    return fec_house_winners_csv(workbook)


impl.fetch = certified_fetch

if __name__ == "__main__":
    impl.main()

    cfg_path = impl.OUT / "cycle.json"
    cfg = json.loads(cfg_path.read_text(encoding="utf-8"))
    cfg.setdefault("scoreability", {})["structural_prior_failures"] = dict(
        sorted(strict.admission.STRUCTURAL_FAILURES.items())
    )
    cfg.setdefault("poll_archive_rules", {})["generic_ballot"] = (
        "Pinned 2022-11-05 pre-election snapshot only; latest admitted poll ended "
        "2022-11-04 and was published by 2022-11-04. No retrospective supplement admitted."
    )
    cfg["poll_archive_rules"]["generic_freshness_rule"] = (
        "Replay uses the frozen model's 35-day generic-ballot window. It does not require "
        "a poll to end on election eve."
    )
    cfg["result_source"] = {
        "authority": "Federal Election Commission",
        "publication": "Federal Elections 2022",
        "sheet": FEC_HOUSE_SHEET,
        "url": FEC_RESULTS_URL,
        "sha256": FEC_RESULTS_SHA256,
        "role": "outcomes only; fetched after predictive inputs and scoreability decisions",
    }
    cfg_path.write_text(json.dumps(cfg, indent=2) + "\n", encoding="utf-8")
