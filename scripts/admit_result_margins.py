#!/usr/bin/env python3
"""Admit certified D/R result margins for Model Lab polling-error evaluation.

Outcome data are fetched only by this explicit admission script. Normal forecast
and validation runs consume the checked-in normalized CSVs and never depend on
these remote sources. Nothing here is a predictive input.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import urllib.request
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
HIST = ROOT / "historical"
UA = "MIDTERM-SIGNALS result-margin admission/1.0 (+https://github.com/MusicOfScience/midterm-signals)"
ALIGNED_INDEPENDENTS = {"ME-S1": "angus king", "VT-S1": "bernie sanders"}
AT_LARGE_2022 = {"AK", "DE", "ND", "SD", "VT", "WY"}
FEC_HOUSE_SHEET = "8. US House Results by State"

STATE_CODES = {
    "ALABAMA":"AL","ALASKA":"AK","ARIZONA":"AZ","ARKANSAS":"AR","CALIFORNIA":"CA",
    "COLORADO":"CO","CONNECTICUT":"CT","DELAWARE":"DE","FLORIDA":"FL","GEORGIA":"GA",
    "HAWAII":"HI","IDAHO":"ID","ILLINOIS":"IL","INDIANA":"IN","IOWA":"IA","KANSAS":"KS",
    "KENTUCKY":"KY","LOUISIANA":"LA","MAINE":"ME","MARYLAND":"MD","MASSACHUSETTS":"MA",
    "MICHIGAN":"MI","MINNESOTA":"MN","MISSISSIPPI":"MS","MISSOURI":"MO","MONTANA":"MT",
    "NEBRASKA":"NE","NEVADA":"NV","NEW HAMPSHIRE":"NH","NEW JERSEY":"NJ","NEW MEXICO":"NM",
    "NEW YORK":"NY","NORTH CAROLINA":"NC","NORTH DAKOTA":"ND","OHIO":"OH","OKLAHOMA":"OK",
    "OREGON":"OR","PENNSYLVANIA":"PA","RHODE ISLAND":"RI","SOUTH CAROLINA":"SC",
    "SOUTH DAKOTA":"SD","TENNESSEE":"TN","TEXAS":"TX","UTAH":"UT","VERMONT":"VT",
    "VIRGINIA":"VA","WASHINGTON":"WA","WEST VIRGINIA":"WV","WISCONSIN":"WI","WYOMING":"WY",
}


def fetch(url: str) -> bytes:
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, timeout=600) as response:
        data = response.read()
    if not data:
        raise RuntimeError(f"empty source: {url}")
    return data


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def number(value) -> float:
    if value in (None, ""):
        return 0.0
    try:
        return float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return 0.0


def norm_name(value) -> str:
    s = str(value or "").lower()
    s = re.sub(r"[“”\"'’`]", "", s)
    s = re.sub(r"\b(jr|sr|ii|iii|iv)\b\.?", "", s)
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return " ".join(s.split())


def party(value) -> str:
    s = str(value or "").strip().upper()
    if s.startswith("DEM") or s in {"D", "DFL", "D/R"} or s.startswith(("D/", "D*")):
        return "DEM"
    if s.startswith("REP") or s == "R" or s.startswith(("R/", "R*")):
        return "REP"
    if s.startswith("IND") or s == "I":
        return "IND"
    return "OTHER"


def state_po(value) -> str:
    s = str(value or "").strip().upper()
    if len(s) == 2 and s.isalpha():
        return s
    return STATE_CODES.get(s, "")


def district_code(state, district) -> str:
    st = state_po(state)
    raw = str(district or "").strip().upper()
    if not st:
        return ""
    if raw in {"", "0", "00", "AL", "AT-LARGE", "AT LARGE", "AT LARGE DISTRICT"}:
        return f"{st}-01"
    m = re.search(r"\d+", raw)
    return f"{st}-{int(m.group(0)):02d}" if m else ""


def existing_winners(year: int) -> dict[str, str]:
    with (HIST / str(year) / "results.csv").open(newline="", encoding="utf-8-sig") as fh:
        return {r["race_code"]: r["winner_party"] for r in csv.DictReader(fh)}


def roster_sides(year: int) -> dict[str, dict[str, str]]:
    """Map pre-election candidate identity to D/R, including already-admitted fallbacks."""
    candidates: dict[str, dict[str, set[str]]] = defaultdict(lambda: defaultdict(set))
    with (HIST / str(year) / "races.csv").open(newline="", encoding="utf-8-sig") as fh:
        for row in csv.DictReader(fh):
            code = (row.get("race_code") or "").strip()
            if not code:
                continue
            for field, side in (("dem_name", "DEM"), ("rep_name", "REP")):
                name = str(row.get(field) or "").strip()
                if not name or name.lower().startswith("(no "):
                    continue
                key = norm_name(name)
                if key:
                    candidates[code][key].add(side)
    cfg = json.loads((HIST / str(year) / "cycle.json").read_text(encoding="utf-8"))
    for row in cfg.get("admission", {}).get("result_party_fallbacks", []):
        code = str(row.get("race_code") or "").strip()
        name = norm_name(row.get("winner_name"))
        side = str(row.get("resolved_side") or "").strip().upper()
        if code and name and side in {"DEM", "REP"}:
            candidates[code][name].add(side)
    out: dict[str, dict[str, str]] = {}
    for code, names in candidates.items():
        out[code] = {name: next(iter(sides)) for name, sides in names.items() if len(sides) == 1}
    return out


def side_for_candidate(code: str, candidate: str, parties: set[str], roster: dict[str, dict[str, str]] | None = None) -> str:
    major = parties & {"DEM", "REP"}
    if len(major) > 1:
        raise RuntimeError(f"candidate appears on both major-party lines: {code} {candidate} {sorted(parties)}")
    if major:
        return next(iter(major))
    aligned = ALIGNED_INDEPENDENTS.get(code)
    if aligned and aligned in norm_name(candidate):
        return "DEM"
    if roster:
        return roster.get(code, {}).get(norm_name(candidate), "OTHER")
    return "OTHER"


def normalized_margin_row(year: int, code: str, candidates: dict[str, dict], vote_key: str = "votes", roster: dict[str, dict[str, str]] | None = None) -> dict:
    totals = {"DEM": 0.0, "REP": 0.0}
    for record in candidates.values():
        side = side_for_candidate(code, record["candidate"], record["parties"], roster)
        if side in totals:
            totals[side] += float(record[vote_key])
    den = totals["DEM"] + totals["REP"]
    if den <= 0:
        raise RuntimeError(f"no D/R result denominator for {year} {code}")
    margin = 100.0 * (totals["DEM"] - totals["REP"]) / den
    winner = "DEM" if margin > 0 else "REP" if margin < 0 else "TIE"
    return {
        "race_code": code,
        "dem_votes": int(round(totals["DEM"])),
        "rep_votes": int(round(totals["REP"])),
        "two_party_votes": int(round(den)),
        "dem_two_party_pct": round(100.0 * totals["DEM"] / den, 6),
        "rep_two_party_pct": round(100.0 * totals["REP"] / den, 6),
        "margin_d_minus_r": round(margin, 6),
        "winner_party_from_margin": winner,
    }


def source_from_cycle_2018(name: str) -> dict:
    cfg = json.loads((HIST / "2018" / "cycle.json").read_text(encoding="utf-8"))
    for item in cfg.get("provenance", []):
        if item.get("name") == name:
            return item
    raise RuntimeError(f"2018 provenance missing {name}")


def admit_2018() -> tuple[list[dict], list[dict]]:
    winners = existing_winners(2018)
    roster = roster_sides(2018)
    groups: dict[str, dict[str, dict]] = defaultdict(dict)
    provenance = []
    for chamber, source_name in (("house", "house_results"), ("senate", "senate_results")):
        src = source_from_cycle_2018(source_name)
        data = fetch(src["url"])
        digest = sha256(data)
        if digest != src["sha256"]:
            raise RuntimeError(f"2018 {source_name} hash changed: {digest} != {src['sha256']}")
        provenance.append({"name": source_name, "url": src["url"], "sha256": digest})
        for row in csv.DictReader(io.StringIO(data.decode("utf-8-sig", errors="replace"))):
            if str(row.get("year") or "") != "2018":
                continue
            if str(row.get("stage") or "").strip().upper() not in {"", "GEN"}:
                continue
            if chamber == "house":
                code = district_code(row.get("state_po") or row.get("state"), row.get("district"))
            else:
                st = state_po(row.get("state_po") or row.get("state"))
                special = str(row.get("special") or "").strip().upper() == "TRUE"
                code = f"{st}-S{2 if special else 1}" if st else ""
            if code not in winners:
                continue
            candidate = str(row.get("candidate") or "").strip()
            key = norm_name(candidate)
            if not key:
                continue
            record = groups[code].setdefault(key, {"candidate": candidate, "parties": set(), "votes": 0.0})
            record["votes"] += number(row.get("candidatevotes"))
            record["parties"].add(party(row.get("party") or row.get("party_simplified")))
    rows = []
    for code in sorted(winners):
        if code not in groups:
            raise RuntimeError(f"2018 certified result rows missing {code}")
        out = normalized_margin_row(2018, code, groups[code], roster=roster)
        if out["winner_party_from_margin"] != winners[code]:
            raise RuntimeError(f"2018 margin winner mismatch {code}: {out['winner_party_from_margin']} != {winners[code]}")
        rows.append(out)
    return rows, provenance


def text(value) -> str:
    return "" if value is None else str(value).strip()


def fec_party_family(value) -> str:
    p = text(value).upper()
    if p in {"D", "DFL", "D/R"} or p.startswith(("D/", "D*", "DEM")):
        return "DEM"
    if p == "R" or p.startswith(("R/", "R*", "REP")):
        return "REP"
    return "OTHER"


def admit_2022() -> tuple[list[dict], list[dict]]:
    cfg = json.loads((HIST / "2022" / "cycle.json").read_text(encoding="utf-8"))
    src = cfg.get("result_source") or {}
    if src.get("authority") != "Federal Election Commission":
        raise RuntimeError(f"2022 result authority is not FEC: {src}")
    data = fetch(src["url"])
    digest = sha256(data)
    if digest != src["sha256"]:
        raise RuntimeError(f"2022 FEC workbook hash changed: {digest} != {src['sha256']}")
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = src.get("sheet") or FEC_HOUSE_SHEET
    if sheet not in wb.sheetnames:
        raise RuntimeError(f"FEC workbook missing {sheet!r}")
    ws = wb[sheet]
    it = ws.iter_rows(values_only=True)
    header = [text(v) for v in next(it)]
    index = {name: i for i, name in enumerate(header) if name}
    required = {"STATE ABBREVIATION", "CANDIDATE NAME", "PARTY", "GENERAL VOTES"}
    missing = sorted(required - set(index))
    district_col = "D" if "D" in index else "DISTRICT" if "DISTRICT" in index else ""
    runoff_cols = [name for name in index if name.startswith("GE RUNOFF ELECTION VOTES")]
    if missing or not district_col or not runoff_cols:
        raise RuntimeError(f"unexpected FEC schema missing={missing} district={district_col!r} runoff={runoff_cols}")
    runoff_col = runoff_cols[0]

    groups: dict[str, dict[str, dict]] = defaultdict(dict)
    for values in it:
        def get(name):
            pos = index[name]
            return values[pos] if pos < len(values) else None
        st = text(get("STATE ABBREVIATION")).upper()
        if st not in STATE_CODES.values():
            continue
        raw_district = text(get(district_col)).upper()
        if "UNEXPIRED" in raw_district:
            continue
        code = district_code(st, raw_district)
        if not code and st in AT_LARGE_2022:
            code = f"{st}-01"
        if not code:
            continue
        candidate = text(get("CANDIDATE NAME"))
        raw_party = text(get("PARTY"))
        key = norm_name(candidate)
        if not key or not raw_party:
            continue
        record = groups[code].setdefault(key, {"candidate": candidate, "parties": set(), "general": 0.0, "runoff": 0.0})
        record["parties"].add(fec_party_family(raw_party))
        record["general"] += number(get("GENERAL VOTES"))
        record["runoff"] += number(get(runoff_col))

    if len(groups) != 435:
        raise RuntimeError(f"FEC margin admission classified {len(groups)} House districts, expected 435")
    winners = existing_winners(2022)
    rows = []
    for code in sorted(groups):
        use_runoff = any(r["runoff"] > 0 for r in groups[code].values())
        out = normalized_margin_row(2022, code, groups[code], "runoff" if use_runoff else "general")
        if code in winners and out["winner_party_from_margin"] != winners[code]:
            raise RuntimeError(f"2022 margin winner mismatch {code}: {out['winner_party_from_margin']} != {winners[code]}")
        out["certified_vote_stage"] = "runoff" if use_runoff else "general"
        rows.append(out)
    return rows, [{"name": "FEC Federal Elections 2022", "url": src["url"], "sha256": digest, "sheet": sheet}]


def write_csv(path: Path, rows: list[dict]) -> str:
    fields = []
    for row in rows:
        for key in row:
            if key not in fields:
                fields.append(key)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    return sha256(path.read_bytes())


def main() -> None:
    rows18, prov18 = admit_2018()
    rows22, prov22 = admit_2022()
    sha18 = write_csv(HIST / "2018" / "result-margins.csv", rows18)
    sha22 = write_csv(HIST / "2022" / "result-margins.csv", rows22)
    manifest = {
        "schema_version": "1.0.0",
        "evaluation_only": True,
        "model_effect": "NONE",
        "rule": "Certified outcome margins may score polling error but may never enter a forecast probability.",
        "cycles": {
            "2018": {"rows": len(rows18), "output": "historical/2018/result-margins.csv", "output_sha256": sha18, "sources": prov18},
            "2022": {"rows": len(rows22), "output": "historical/2022/result-margins.csv", "output_sha256": sha22, "sources": prov22},
        },
        "transformation": "Candidate/fusion lines are collapsed within race; Democratic and Republican certified votes are summed; margin is 100*(D-R)/(D+R). A 2018 result line without a usable party label may inherit D/R only from an unambiguous normalized candidate match in the pre-election roster or from an identity fallback already admitted and recorded in the 2018 cycle manifest. Angus King and Bernie Sanders retain the model's predeclared Democratic-aligned control side. 2022 uses certified runoff votes when the FEC sheet reports a House runoff, otherwise certified general-election votes.",
    }
    (HIST / "result-margin-provenance.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"2018_rows": len(rows18), "2022_rows": len(rows22), "2018_sha": sha18, "2022_sha": sha22}, indent=2))


if __name__ == "__main__":
    main()
